"""Tests for Enterprise DOCX Report Generator."""

import os
import tempfile
import pytest
from unittest.mock import patch

from src.client.report_state import ReportState, Finding, ToolStep
from src.utils.report_generator import generate_docx_report


@pytest.fixture
def sample_report_data():
    """Create a sample report data dict (simulating ReportState.to_dict())."""
    state = ReportState(target="http://testphp.vulnweb.com", scan_mode="recon")

    # Add findings
    state.add_finding(Finding(
        title="Apache 2.4.49 Path Traversal (CVE-2021-41773)",
        severity="CRITICAL", description="Path traversal allows arbitrary file read",
        impact="Complete server compromise", remediation="Upgrade Apache to 2.4.51+",
        tool_source="docker_nuclei_scan", raw_evidence="CVE-2021-41773 detected",
    ))
    state.add_finding(Finding(
        title="MySQL Port 3306 Exposed",
        severity="MEDIUM", description="MySQL accessible from external network",
        impact="Database data leak", remediation="Restrict to localhost only",
        tool_source="docker_scan_ports_deep", raw_evidence="3306/tcp open mysql",
    ))
    state.add_finding(Finding(
        title="Missing Security Headers",
        severity="LOW", description="X-Frame-Options and CSP headers missing",
        impact="Clickjacking risk", remediation="Add security headers",
        tool_source="docker_whatweb", raw_evidence="No X-Frame-Options",
    ))

    # Add methodology
    state.add_step(ToolStep(
        step_number=1, tool_name="docker_scan_ports_fast",
        arguments={"target": "testphp.vulnweb.com"},
        result_snippet="22,80,443,3306", status="SUCCESS",
        reporter_comment="4 ports discovered",
    ))
    state.add_step(ToolStep(
        step_number=2, tool_name="docker_nuclei_scan",
        arguments={"target": "http://testphp.vulnweb.com"},
        result_snippet="CVE-2021-41773 found", status="SUCCESS",
        reporter_comment="Critical CVE detected!",
    ))

    state.update_risk_score(8.5)
    state.executive_summary = "Cuộc đánh giá phát hiện lỗ hổng nghiêm trọng Apache CVE-2021-41773."
    state.conclusion = "Hệ thống có rủi ro CAO. Cần khắc phục ngay."
    state.add_recommendation("Nâng cấp Apache lên 2.4.51+")
    state.add_recommendation("Chặn port 3306 từ bên ngoài")

    # Populate attack surface
    state.attack_surface.open_ports = {
        80: {"service": "http", "deep_scanned": True},
        3306: {"service": "mysql", "deep_scanned": False},
    }
    state.attack_surface.detected_technologies = ["Apache 2.4.49", "PHP 7.4"]
    state.attack_surface.parameterized_endpoints = {
        "http://testphp.vulnweb.com/listproducts.php?cat=1": {"tested_sqli": True}
    }
    state.attack_surface.login_forms = [{"url": "http://testphp.vulnweb.com/login.php", "tested_bruteforce": False}]
    state.attack_surface.subdomains = ["dev.vulnweb.com", "admin.vulnweb.com"]

    state.finalize(red_teamer_answer="Final analysis complete.")

    return state.to_dict()


@pytest.fixture
def empty_report_data():
    """Create empty report data (no findings)."""
    state = ReportState(target="192.168.1.1", scan_mode="full")
    state.finalize()
    return state.to_dict()


class TestDocxGeneration:
    """Test DOCX report generation."""

    @patch("src.utils.report_generator.cfg_get")
    def test_generates_docx_file(self, mock_cfg_get, sample_report_data, tmp_path):
        """Should generate a .docx file that exists."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Test Corp",
            "reports.classification": "INTERNAL",
        }.get(key, default)

        result = generate_docx_report(sample_report_data)
        assert result.endswith(".docx")
        assert os.path.exists(result)

    @patch("src.utils.report_generator.cfg_get")
    def test_docx_file_not_empty(self, mock_cfg_get, sample_report_data, tmp_path):
        """Generated DOCX should not be empty."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Test Corp",
            "reports.classification": "CONFIDENTIAL",
        }.get(key, default)

        result = generate_docx_report(sample_report_data)
        assert os.path.getsize(result) > 1000  # At least 1KB

    @patch("src.utils.report_generator.cfg_get")
    def test_empty_report_generates_ok(self, mock_cfg_get, empty_report_data, tmp_path):
        """Should handle empty report data without crashing."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Empty Test",
            "reports.classification": "INTERNAL",
        }.get(key, default)

        result = generate_docx_report(empty_report_data)
        assert result.endswith(".docx")
        assert os.path.exists(result)

    @patch("src.utils.report_generator.cfg_get")
    def test_docx_contains_expected_content(self, mock_cfg_get, sample_report_data, tmp_path):
        """DOCX should contain the target, findings, and risk info."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Content Test Corp",
            "reports.classification": "SECRET",
        }.get(key, default)

        from docx import Document as DocxDocument

        result = generate_docx_report(sample_report_data)
        doc = DocxDocument(result)

        # Extract all text from paragraphs AND table cells
        full_text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text

        # Check key content
        assert "testphp.vulnweb.com" in full_text
        assert "Content Test Corp" in full_text
        assert "CRITICAL" in full_text
        assert "CVE-2021-41773" in full_text

    @patch("src.utils.report_generator.cfg_get")
    def test_docx_has_tables(self, mock_cfg_get, sample_report_data, tmp_path):
        """DOCX should contain tables (risk summary, methodology, findings)."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Table Test",
            "reports.classification": "INTERNAL",
        }.get(key, default)

        from docx import Document as DocxDocument

        result = generate_docx_report(sample_report_data)
        doc = DocxDocument(result)

        # Should have at least 3 tables: cover info, risk summary, methodology, findings detail, risk matrix
        assert len(doc.tables) >= 3

    @patch("src.utils.report_generator.cfg_get")
    def test_docx_table_header_and_cant_split(self, mock_cfg_get, sample_report_data, tmp_path):
        """DOCX tables should have tblHeader on headers and cantSplit on all rows."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Table Split Test",
            "reports.classification": "INTERNAL",
        }.get(key, default)

        from docx import Document as DocxDocument

        result = generate_docx_report(sample_report_data)
        doc = DocxDocument(result)

        tbl_header_count = 0
        cant_split_count = 0

        for table in doc.tables:
            for row in table.rows:
                trPr = row._tr.trPr
                if trPr is not None:
                    if trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblHeader") is not None:
                        tbl_header_count += 1
                    if trPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is not None:
                        cant_split_count += 1

        # Must have at least 3 table headers configured (risk summary, methodology, risk matrix)
        assert tbl_header_count >= 3
        # Must have rows with cantSplit to prevent splitting across pages
        assert cant_split_count >= 10

    @patch("src.utils.report_generator.cfg_get")
    def test_docx_contains_attack_surface(self, mock_cfg_get, sample_report_data, tmp_path):
        """DOCX should contain the Attack Surface section, open ports, and technologies."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Attack Surface Corp",
            "reports.classification": "CONFIDENTIAL",
        }.get(key, default)

        from docx import Document as DocxDocument

        result = generate_docx_report(sample_report_data)
        doc = DocxDocument(result)

        full_text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text

        assert "Bề mặt Tấn công" in full_text
        assert "80/tcp" in full_text
        assert "3306/tcp" in full_text
        assert "Apache 2.4.49" in full_text
        assert "dev.vulnweb.com" in full_text

    @patch("src.utils.report_generator.cfg_get")
    def test_docx_findings_contain_cve_and_cvss(self, mock_cfg_get, sample_report_data, tmp_path):
        """Finding details table in DOCX should render CVE ID and CVSS if present."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "CVE Test Corp",
            "reports.classification": "CONFIDENTIAL",
        }.get(key, default)

        sample_report_data["findings"][0]["cve_id"] = "CVE-2021-41773"
        sample_report_data["findings"][0]["cvss_score"] = 9.8

        from docx import Document as DocxDocument

        result = generate_docx_report(sample_report_data)
        doc = DocxDocument(result)

        full_text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text

        assert "Mã CVE:" in full_text
        assert "CVE-2021-41773" in full_text
        assert "Điểm CVSS:" in full_text
        assert "9.8" in full_text

    @patch("src.utils.report_generator.cfg_get")
    def test_empty_report_uses_v4_1_text(self, mock_cfg_get, empty_report_data, tmp_path):
        """Default executive summary in empty report should mention MLSecOps Agent v4.1."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Version Test",
            "reports.classification": "INTERNAL",
        }.get(key, default)

        from docx import Document as DocxDocument

        result = generate_docx_report(empty_report_data)
        doc = DocxDocument(result)

        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "MLSecOps Agent v4.1" in full_text

    @patch("src.utils.report_generator.cfg_get")
    def test_docx_contains_rag_memory_section(self, mock_cfg_get, sample_report_data, tmp_path):
        """DOCX should include RAG long-term memory section and tables."""
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "RAG Test Corp",
            "reports.classification": "CONFIDENTIAL",
        }.get(key, default)

        from docx import Document as DocxDocument

        sample_report_data["rag_applied_patterns"] = [
            {"id": "ap_old_1", "cwe_id": "CWE-89", "entry_point": "http://target/api", "successful_vector": "sqlmap --batch", "similarity_score": 0.95}
        ]
        sample_report_data["rag_stored_patterns"] = [
            {"id": "ap_new_1", "cwe_id": "CWE-79", "entry_point": "/search", "successful_vector": "<svg/onload=alert(1)>", "agent_reasoning": "XSS bypass success"}
        ]

        result = generate_docx_report(sample_report_data)
        doc = DocxDocument(result)

        full_text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text

        assert "Trí nhớ Chiến thuật Dài hạn & Khả năng Tự học" in full_text
        assert "CWE-89" in full_text
        assert "sqlmap --batch" in full_text
        assert "CWE-79" in full_text

    def test_cell_shading_and_border_caching(self):
        """Verify _set_cell_shading and _set_cell_border with template caching."""
        from docx import Document as DocxDocument
        from src.utils.report_generator import _set_cell_shading, _set_cell_border, _SHADING_CACHE

        doc = DocxDocument()
        table = doc.add_table(rows=2, cols=2)
        cell1 = table.cell(0, 0)
        cell2 = table.cell(0, 1)

        # First call caches "1E3A5F"
        _set_cell_shading(cell1, "1E3A5F")
        assert "1E3A5F" in _SHADING_CACHE

        # Second call reuses deepcopy
        _set_cell_shading(cell2, "1E3A5F")
        assert cell1._tc.get_or_add_tcPr().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd") is not None
        assert cell2._tc.get_or_add_tcPr().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd") is not None

        # Border test
        _set_cell_border(cell1, bottom={"sz": "6", "val": "single", "color": "1E3A5F"})
        assert cell1._tc.get_or_add_tcPr().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcBorders") is not None


