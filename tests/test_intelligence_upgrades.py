"""Unit tests for Supreme Dual-AI Intelligence and Attack Surface Reporting Upgrades.

Covers:
- Dynamic attack surface expansion & stagnation tracking logic.
- Tactical roadmap and progressive goal calculation.
- SOC Analyst objective assessment & blocking factor feedback formatting.
- Targeted mission objective guard verification.
- Enterprise Word DOCX report generation with full attack surface (exposed files, CORS, alive subdomains).
"""

import json
import os
import pathlib
import pytest
import tempfile
from docx import Document

from src.client.report_state import ReportState, AttackSurfaceGraph, Finding, ToolStep
from src.utils.report_generator import generate_docx_report
from src.client.orchestrator import _evaluate_mission_guard


class TestAttackSurfaceTrackingAndStagnation:
    """Tests for dynamic attack surface discovery and stagnation tracking logic."""

    def test_surface_expansion_tracking(self):
        state = ReportState(target="http://test.local", scan_mode="full", mission_objective="Audit full target")
        assert not state.attack_surface.has_attack_surface()

        # Update with port scan result
        port_result = json.dumps({"open_ports": [{"port": 80, "service": "http"}, {"port": 22, "service": "ssh"}]})
        state.update_attack_surface("docker_scan_ports_fast", {"target": "test.local"}, port_result)
        assert 80 in state.attack_surface.open_ports
        assert 22 in state.attack_surface.open_ports
        assert state.attack_surface.has_attack_surface()

        # Update with crawling result
        crawl_result = json.dumps({"endpoints": ["http://test.local/search?q=1", "http://test.local/login"]})
        state.update_attack_surface("docker_crawl_web", {"url": "http://test.local"}, crawl_result)
        assert "http://test.local/search?q=1" in state.attack_surface.parameterized_endpoints
        assert any(f["url"] == "http://test.local/login" for f in state.attack_surface.login_forms)

        # Update with sensitive files scan
        files_result = json.dumps({"exposed_files": [{"path": "/.env", "risk_type": "Environment Configuration", "severity": "HIGH", "status_code": 200}]})
        state.update_attack_surface("docker_sensitive_files_scan", {"target_url": "http://test.local"}, files_result)
        assert len(state.attack_surface.exposed_sensitive_files) == 1
        assert state.attack_surface.exposed_sensitive_files[0]["path"] == "/.env"

        # Update with CORS scan
        cors_result = json.dumps({"cors_misconfigurations": [{"url": "http://test.local/api", "type": "Reflected Origin", "severity": "MEDIUM"}]})
        state.update_attack_surface("docker_cors_scan", {"target_url": "http://test.local"}, cors_result)
        assert len(state.attack_surface.cors_issues) == 1

        # Update with HTTPX probe
        httpx_result = json.dumps({"alive_targets": [{"url": "http://admin.test.local", "alive": True, "status_code": 200, "title": "Dashboard", "webserver": "nginx"}]})
        state.update_attack_surface("docker_httpx_probe", {"target": "admin.test.local"}, httpx_result)
        assert len(state.attack_surface.alive_subdomains) == 1
        assert state.attack_surface.alive_subdomains[0]["url"] == "http://admin.test.local"

    def test_surface_count_calculation(self):
        """Verify the compound attack surface count used by the orchestrator to detect progress."""
        state = ReportState(target="http://test.local")
        state.attack_surface.open_ports = {80: {}, 443: {}}
        state.attack_surface.parameterized_endpoints = {"http://test.local/view?id=1": {}}
        state.attack_surface.subdomains = ["api.test.local", "dev.test.local"]
        state.attack_surface.alive_subdomains = [{"url": "http://api.test.local"}]
        state.attack_surface.exposed_sensitive_files = [{"path": "/.git/config"}]
        state.attack_surface.cors_issues = [{"url": "http://test.local"}]

        surface_count = (
            len(state.attack_surface.open_ports)
            + len(state.attack_surface.parameterized_endpoints)
            + len(state.attack_surface.subdomains)
            + len(state.attack_surface.alive_subdomains)
            + len(state.attack_surface.exposed_sensitive_files)
            + len(state.attack_surface.cors_issues)
        )
        assert surface_count == 2 + 1 + 2 + 1 + 1 + 1  # 8 items


class TestTacticalRoadmapAndProgress:
    """Tests for progressive goal calculation and prioritized tactical roadmap."""

    def test_calculate_goal_progress_increases_with_discovery(self):
        state = ReportState(target="http://test.local", scan_mode="full")
        p0 = state.calculate_goal_progress()
        assert p0 == 5

        # Execute port scan step
        state.add_step(ToolStep(step_number=1, tool_name="docker_scan_ports_fast", arguments={"target": "test.local"}, result_snippet="ports", status="SUCCESS"))
        p1 = state.calculate_goal_progress()
        assert p1 > p0  # Phase 1: 10 + 20 = 30

        # Execute web crawl step
        state.add_step(ToolStep(step_number=2, tool_name="docker_crawl_web", arguments={"url": "http://test.local"}, result_snippet="urls", status="SUCCESS"))
        p2 = state.calculate_goal_progress()
        assert p2 > p1  # Phase 2: 30 + 25 = 55

        # Execute vuln scan step + add finding
        state.add_step(ToolStep(step_number=3, tool_name="docker_nuclei_scan", arguments={"target": "http://test.local"}, result_snippet="vulns", status="SUCCESS"))
        state.add_finding(Finding(title="Exposed Config", severity="HIGH"))
        p3 = state.calculate_goal_progress()
        assert p3 > p2  # Phase 3 + High finding: 55 + 25 + 15 = 95

    def test_tactical_roadmap_priorities(self):
        state = ReportState(target="http://test.local", scan_mode="full")
        state.attack_surface.open_ports = {80: {}, 22: {}}
        state.attack_surface.parameterized_endpoints = {"http://test.local/?id=1": {"tested_sqli": False, "tested_xss": False}}
        state.attack_surface.subdomains = ["sub1.test.local"]

        roadmap = state.get_tactical_roadmap()
        assert "untested_actions" in roadmap
        assert "top_actions" in roadmap
        assert len(roadmap["untested_actions"]) > 0

        # Top priority should be SQLi or Alive probe for unprobed subdomains
        tools_in_actions = [a["tool"] for a in roadmap["untested_actions"]]
        assert "docker_sqlmap_scan" in tools_in_actions
        assert "docker_httpx_probe" in tools_in_actions


class TestReporterFeedbackBlockFormatting:
    """Tests for SOC Analyst objective assessment and feedback block injection."""

    def test_feedback_block_with_objective_and_blocking_factor(self):
        state = ReportState(target="http://test.local", mission_objective="Tìm lỗ hổng SQLi")
        state.risk_score = 7.5

        feedback = state.get_reporter_feedback_block(
            latest_suggestion="Hãy khai thác endpoint /search?q= với SQLMap",
            latest_comment="Phát hiện endpoint nghi vấn SQLi",
            latest_findings=[Finding(title="Potential SQL Injection", severity="HIGH")],
            objective_assessment="Đã xác định endpoint chứa tham số, đang tiến hành injection",
            blocking_factor="WAF Cloudflare phản hồi 403",
        )

        assert "SOC ANALYST" in feedback
        assert "Tiến độ:" in feedback
        assert "Risk Score: 7.5/10" in feedback
        assert "🎯 Đánh giá mục tiêu: Đã xác định endpoint chứa tham số, đang tiến hành injection | Cản trở: WAF Cloudflare phản hồi 403" in feedback
        assert "HIGH" in feedback
        assert "Potential SQL Injection" in feedback
        assert "Gợi ý: Hãy khai thác endpoint /search?q= với SQLMap" in feedback

    def test_feedback_block_backward_compatible_empty_objective(self):
        state = ReportState(target="http://test.local")
        state.risk_score = 3.0

        feedback = state.get_reporter_feedback_block(
            latest_suggestion="Quét cổng dịch vụ",
            latest_comment="Bước bắt đầu",
            latest_findings=[],
        )
        assert "🎯 Đánh giá mục tiêu" not in feedback
        assert "Risk Score: 3.0/10" in feedback


class TestDOCXAttackSurfaceRendering:
    """Tests for rendering the full attack surface in the Enterprise Word DOCX report."""

    def test_docx_contains_all_attack_surface_tables(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            out_path = os.path.join(tmpdir, "test_report.docx")

            report_data = {
                "target": "https://target.corp",
                "company_name": "Pentest Corp",
                "classification": "CONFIDENTIAL",
                "scan_mode": "full",
                "start_time": "2026-09-18 10:00:00",
                "end_time": "2026-09-18 10:30:00",
                "risk_score": 8.0,
                "risk_label": "HIGH",
                "severity_counts": {"CRITICAL": 0, "HIGH": 2, "MEDIUM": 1, "LOW": 0, "INFO": 1},
                "executive_summary": "Kiểm thử thâm nhập toàn diện hoàn thành xuất sắc.",
                "methodology": [
                    {"step_number": 1, "tool_name": "docker_scan_ports_fast", "target": "target.corp", "status": "SUCCESS", "timestamp": "10:01:00", "reporter_comment": "2 cổng mở"}
                ],
                "findings": [
                    {"title": "Exposed .env Configuration", "severity": "HIGH", "description": "Lộ tệp cấu hình chứa khóa bí mật.", "evidence": "DB_PASSWORD=secret", "recommendation": "Xóa hoặc chặn truy cập."}
                ],
                "recommendations": ["Chặn tệp nhạy cảm trên web server."],
                "attack_surface": {
                    "open_ports": {80: {"service": "http", "deep_scanned": True}, 443: {"service": "https", "deep_scanned": True}},
                    "parameterized_endpoints": {
                        "https://target.corp/api/user?id=1": {"tested_sqli": True, "tested_xss": True},
                        "https://target.corp/search?q=test": {"tested_sqli": False, "tested_xss": False},
                    },
                    "login_forms": [{"url": "https://target.corp/login", "tested_bruteforce": True}],
                    "detected_technologies": ["Nginx", "PHP 8.2", "Laravel"],
                    "subdomains": ["api.target.corp", "admin.target.corp"],
                    "alive_subdomains": [
                        {"url": "https://api.target.corp", "status_code": 200, "title": "API Gateway", "webserver": "nginx/1.24"}
                    ],
                    "exposed_sensitive_files": [
                        {"path": "/.env", "risk_type": "Environment Configuration", "severity": "HIGH", "status_code": 200},
                        {"path": "/.git/HEAD", "risk_type": "Git Repository Exposure", "severity": "MEDIUM", "status_code": 200}
                    ],
                    "cors_issues": [
                        {"url": "https://target.corp/api/data", "type": "Origin Reflection with Credentials", "severity": "HIGH", "description": "Access-Control-Allow-Origin reflects origin with credentials"}
                    ]
                }
            }

            result_path = generate_docx_report(report_data, out_path)
            assert os.path.exists(result_path)
            assert os.path.getsize(result_path) > 5000

            doc = Document(result_path)
            all_text = " ".join(p.text for p in doc.paragraphs)
            all_cell_text = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            combined_text = all_text + " " + all_cell_text

            # Check for attack surface headings and table contents
            assert "Phân tích Bề mặt Tấn công" in combined_text
            assert "Các Cổng Dịch Vụ Mở" in combined_text
            assert "Tập Tin & Đường Dẫn Nhạy Cảm Lộ Lọt" in combined_text
            assert "Cấu Hình CORS & Tiêu Đề Bảo Mật" in combined_text
            assert "Tên Miền Phụ Đang Hoạt Động" in combined_text
            assert "/.env" in combined_text
            assert "Git Repository Exposure" in combined_text
            assert "Origin Reflection with Credentials" in combined_text
            assert "https://api.target.corp" in combined_text
            assert "API Gateway" in combined_text
            assert "SQLi: Đã test | XSS: Đã test" in combined_text
            assert "SQLi: Chưa test | XSS: Chưa test" in combined_text


class TestMissionGuardVerification:
    """Tests for _evaluate_mission_guard in orchestrator."""

    def test_guard_flags_untested_sqli_when_requested(self):
        state = ReportState(target="http://test.local", mission_objective="Tìm lỗ hổng SQL Injection và trích xuất dữ liệu")
        state.attack_surface.parameterized_endpoints = {"http://test.local/item?id=1": {"tested_sqli": False}}

        should_guard, unfulfilled = _evaluate_mission_guard(
            report_state=state,
            tools_called={"docker_crawl_web"},
            progress_pct=40,
            untapped_actions=[{"action": "docker_sqlmap_scan", "priority": "CRITICAL"}],
            iteration=3,
            max_iterations=15,
            guard_attempts=0,
        )
        assert should_guard is True
        assert any("SQL Injection" in u for u in unfulfilled)

    def test_guard_flags_untested_xss_when_requested(self):
        state = ReportState(target="http://test.local", mission_objective="Kiểm tra XSS lỗ hổng web")
        state.attack_surface.parameterized_endpoints = {"http://test.local/search?q=a": {"tested_xss": False}}

        should_guard, unfulfilled = _evaluate_mission_guard(
            report_state=state,
            tools_called={"docker_crawl_web"},
            progress_pct=40,
            untapped_actions=[{"action": "docker_xss_scan", "priority": "HIGH"}],
            iteration=3,
            max_iterations=15,
            guard_attempts=0,
        )
        assert should_guard is True
        assert any("XSS" in u for u in unfulfilled)

    def test_guard_flags_untested_sensitive_files(self):
        state = ReportState(target="http://test.local", mission_objective="Quét các file nhạy cảm .env backup")

        should_guard, unfulfilled = _evaluate_mission_guard(
            report_state=state,
            tools_called={"docker_scan_ports_fast"},
            progress_pct=30,
            untapped_actions=[],
            iteration=2,
            max_iterations=15,
            guard_attempts=0,
        )
        assert should_guard is True
        assert any("tệp tin nhạy cảm" in u for u in unfulfilled)

    def test_guard_allows_final_answer_when_requirements_met(self):
        state = ReportState(target="http://test.local", mission_objective="Kiểm tra SQLi")
        state.attack_surface.parameterized_endpoints = {"http://test.local/item?id=1": {"tested_sqli": True}}

        should_guard, unfulfilled = _evaluate_mission_guard(
            report_state=state,
            tools_called={"docker_crawl_web", "docker_sqlmap_scan"},
            progress_pct=85,
            untapped_actions=[],
            iteration=5,
            max_iterations=15,
            guard_attempts=0,
        )
        assert should_guard is False
        assert len(unfulfilled) == 0

    def test_guard_bypassed_after_max_attempts(self):
        state = ReportState(target="http://test.local", mission_objective="Tìm SQLi")
        state.attack_surface.parameterized_endpoints = {"http://test.local/item?id=1": {"tested_sqli": False}}

        should_guard, unfulfilled = _evaluate_mission_guard(
            report_state=state,
            tools_called=set(),
            progress_pct=20,
            untapped_actions=[{"action": "docker_sqlmap_scan", "priority": "CRITICAL"}],
            iteration=5,
            max_iterations=15,
            guard_attempts=3,  # Already challenged 3 times
        )
        assert should_guard is False

    def test_guard_bypassed_near_iteration_limit(self):
        state = ReportState(target="http://test.local", mission_objective="Tìm SQLi")
        state.attack_surface.parameterized_endpoints = {"http://test.local/item?id=1": {"tested_sqli": False}}

        should_guard, unfulfilled = _evaluate_mission_guard(
            report_state=state,
            tools_called=set(),
            progress_pct=20,
            untapped_actions=[{"action": "docker_sqlmap_scan", "priority": "CRITICAL"}],
            iteration=13,  # iteration >= 15 - 2
            max_iterations=15,
            guard_attempts=0,
        )
        assert should_guard is False
