"""Tests for All-Encompassing Full-Spectrum Reporting Engine.

Verifies that DOCX, Markdown, and JSON outputs report every aspect
of reconnaissance, vulnerability assessment, exploitation, attack chaining,
kill chain timeline, prioritized remediation roadmap, and technical appendix.
"""

import os
import tempfile
import pytest
from unittest.mock import patch
from docx import Document as DocxDocument

from src.client.report_state import ReportState, Finding, ToolStep, ObjectiveMilestone
from src.utils.report_generator import generate_docx_report


@pytest.fixture
def comprehensive_report_state():
    """Create a fully-populated ReportState with all attack surfaces, compound threats, and milestones."""
    state = ReportState(
        target="http://testphp.vulnweb.com:8080/admin/login?debug=1",
        scan_mode="full",
        mission_objective="Thẩm định toàn diện hệ thống web quản trị và cơ sở dữ liệu",
    )

    # 1. Milestones
    state.milestones = [
        ObjectiveMilestone("recon", "Trinh Sát Bề Mặt", "Rà quét cổng và thu thập URL", "COMPLETED", 100, ["docker_scan_ports_fast", "docker_crawl_web"], 2),
        ObjectiveMilestone("vuln", "Rà Soát Lỗ Hổng", "Quét CVE và kiểm thử SQLi/XSS", "COMPLETED", 100, ["docker_nuclei_scan", "docker_xss_scan"], 3),
        ObjectiveMilestone("exploit", "Xác Minh Khai Thác", "Xác minh khả năng bypass auth", "IN_PROGRESS", 60, ["docker_sqlmap_scan"], 1),
        ObjectiveMilestone("report", "Báo Cáo Toàn Diện", "Tổng hợp và đề xuất lộ trình", "COMPLETED", 100, [], 0),
    ]

    # 2. Findings
    state.add_finding(Finding(
        title="Apache 2.4.49 Remote Code Execution (CVE-2021-41773)",
        severity="CRITICAL",
        description="Path traversal and RCE through cgi-bin directory",
        impact="Toàn quyền chiếm quyền điều khiển máy chủ",
        remediation="Nâng cấp Apache lên phiên bản 2.4.51 trở lên",
        tool_source="docker_nuclei_scan",
        raw_evidence="POST /cgi-bin/.%2e/.%2e/.%2e/.%2e/bin/sh -> uid=0(root)",
        cve_id="CVE-2021-41773",
        cvss_score=9.8,
    ))
    state.add_finding(Finding(
        title="SQL Injection trên tham số cat",
        severity="HIGH",
        description="Boolean-based blind SQL Injection",
        impact="Rò rỉ toàn bộ cơ sở dữ liệu khách hàng",
        remediation="Sử dụng Parameterized Queries hoặc ORM",
        tool_source="docker_sqlmap_scan",
        raw_evidence="Parameter 'cat' is vulnerable: AND 1=1 vs AND 1=2",
        cve_id="",
        cvss_score=8.5,
    ))
    state.add_finding(Finding(
        title="Cookie Quản trị viên thiếu cờ HttpOnly",
        severity="MEDIUM",
        description="Cookie phiên có thể bị đọc trộm qua mã độc JavaScript",
        impact="Nguy cơ chiếm quyền phiên đăng nhập (Session Hijacking)",
        remediation="Bổ sung cờ HttpOnly; Secure; SameSite=Strict",
        tool_source="docker_cookie_security_scan",
        raw_evidence="Set-Cookie: admin_session=xyz123; path=/",
    ))

    # Trigger compound correlation
    state.correlate_compound_threats()

    # 3. Attack Surface
    state.attack_surface.open_ports = {
        80: {"service": "http", "deep_scanned": True},
        8080: {"service": "http-proxy", "deep_scanned": True},
        3306: {"service": "mysql", "deep_scanned": False},
    }
    state.attack_surface.detected_technologies = ["Apache 2.4.49", "PHP 7.4", "MySQL 5.7"]
    state.attack_surface.parameterized_endpoints = {
        "http://testphp.vulnweb.com/listproducts.php?cat=1": {"tested_sqli": True, "tested_xss": True},
        "http://testphp.vulnweb.com/artists.php?artist=2": {"tested_sqli": False, "tested_xss": True},
    }
    state.attack_surface.login_forms = [
        {"url": "http://testphp.vulnweb.com/login.php", "tested_bruteforce": True}
    ]
    state.attack_surface.exposed_sensitive_files = [
        {"path": "/.env", "risk_type": "Credential Leak", "severity": "HIGH", "status_code": 200},
        {"path": "/backup.sql", "risk_type": "Database Dump", "severity": "HIGH", "status_code": 200},
    ]
    state.attack_surface.cors_issues = [
        {"url": "http://testphp.vulnweb.com/api/user", "type": "Reflected Origin", "severity": "MEDIUM", "description": "Access-Control-Allow-Origin: *"}
    ]
    state.attack_surface.alive_subdomains = [
        {"url": "http://admin.vulnweb.com", "status_code": 200, "title": "Admin Portal", "webserver": "Apache"}
    ]
    state.attack_surface.ssl_cert_info = {
        "subject_cn": "vulnweb.com",
        "issuer_org": "Let's Encrypt",
        "valid_to": "2027-01-01",
        "days_until_expiry": 300,
        "tls_version": "TLSv1.3",
        "key_bits": 2048,
        "subject_alternative_names": ["vulnweb.com", "www.vulnweb.com"],
        "issues": ["Thiếu HSTS header"],
    }
    state.attack_surface.dns_security_info = {
        "spf": {"status": "valid", "record": "v=spf1 include:_spf.google.com ~all"},
        "dmarc": {"status": "missing", "policy": "none"},
        "dnssec": {"dnssec_enabled": False},
        "security_issues": ["Thiếu chính sách DMARC reject", "Chưa kích hoạt DNSSEC"],
    }
    state.attack_surface.hidden_discovered_paths = [
        "/admin/config.php",
        "/secret/debug_log.txt",
    ]
    state.attack_surface.cookie_issues = [
        {"name": "admin_session", "missing_flags": ["HttpOnly", "Secure", "SameSite"], "severity": "MEDIUM", "risk": "Đánh cắp phiên qua XSS"}
    ]

    # 4. Methodology
    state.add_step(ToolStep(
        step_number=1,
        tool_name="docker_scan_ports_fast",
        arguments={"target": "testphp.vulnweb.com"},
        result_snippet="Ports 80, 8080, 3306 open",
        status="SUCCESS",
        duration_seconds=2.1,
        reporter_comment="Phát hiện 3 cổng mở bao gồm cổng MySQL 3306 công khai",
    ))
    state.add_step(ToolStep(
        step_number=2,
        tool_name="docker_nuclei_scan",
        arguments={"target": "http://testphp.vulnweb.com:8080"},
        result_snippet="CVE-2021-41773 Apache Path Traversal confirmed",
        status="SUCCESS",
        duration_seconds=5.4,
        reporter_comment="Xác nhận lỗ hổng thực thi mã từ xa RCE Apache",
    ))

    state.update_risk_score(9.8)
    state.executive_summary = "Đánh giá an toàn thông tin phát hiện nguy cơ chiếm đoạt máy chủ hoàn toàn."
    state.conclusion = "Hệ thống tồn tại rủi ro CỰC KỲ NGHIÊM TRỌNG. Yêu cầu vá khẩn cấp trong 24 giờ."
    state.finalize(red_teamer_answer="Đã khai thác thành công RCE và kiểm chứng đường dẫn tấn công.")

    return state


class TestAllEncompassingDocxReport:
    """Verify DOCX generation includes all extended sections."""

    @patch("src.utils.report_generator.cfg_get")
    def test_docx_contains_all_extended_sections(self, mock_cfg_get, comprehensive_report_state, tmp_path):
        mock_cfg_get.side_effect = lambda key, default=None: {
            "reports.output_dir": str(tmp_path),
            "reports.company_name": "Cyber Defense Unit",
            "reports.classification": "TOP SECRET",
        }.get(key, default)

        report_data = comprehensive_report_state.to_dict()
        docx_path = generate_docx_report(report_data)

        assert os.path.exists(docx_path)
        assert os.path.getsize(docx_path) > 5000

        doc = DocxDocument(docx_path)

        full_text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    full_text += "\n" + cell.text

        # 1. Scope & Milestones
        assert "Phạm vi Đánh giá & Cột mốc Sứ mệnh" in full_text
        assert "Trinh Sát Bề Mặt" in full_text
        assert "Rà Soát Lỗ Hổng" in full_text
        assert "Xác Minh Khai Thác" in full_text

        # 2. Attack Surface Subsections
        assert "Các Cổng Dịch Vụ Mở" in full_text
        assert "8080/tcp" in full_text
        assert "Tập Tin & Đường Dẫn Nhạy Cảm Lộ Lọt" in full_text
        assert "backup.sql" in full_text
        assert "An toàn Hạ tầng DNS & Chống Giả mạo Email" in full_text
        assert "An Toàn Cookie & Quản Lý Phiên" in full_text

        # 3. Compound Attack Scenarios
        assert "Kịch Bản Khai Thác Phức Hợp" in full_text

        # 4. Detailed Findings
        assert "CVE-2021-41773" in full_text
        assert "9.8" in full_text

        # 5. Prioritized Remediation Roadmap
        assert "Lộ trình Khắc phục Phân tầng Thời gian" in full_text
        assert "Khẩn cấp (< 24 giờ)" in full_text
        assert "Ngắn hạn (< 7 ngày)" in full_text
        assert "Trung hạn (< 30 ngày)" in full_text

        # 6. Technical Execution Appendix
        assert "Phụ lục Kỹ thuật Thực thi" in full_text
        assert "docker_nuclei_scan" in full_text
        assert "docker_scan_ports_fast" in full_text


class TestAllEncompassingMarkdownReport:
    """Verify Markdown generation contains all complete tables and sections."""

    def test_markdown_contains_all_comprehensive_tables(self, comprehensive_report_state):
        md_text = comprehensive_report_state.to_markdown()

        # 1. Title & Metadata
        assert "# BÁO CÁO ĐÁNH GIÁ AN TOÀN THÔNG TIN TOÀN DIỆN" in md_text
        assert "Mục tiêu yêu cầu" in md_text
        assert "Chế độ kiểm thử" in md_text
        assert "Chỉ số Rủi ro Tổng thể" in md_text
        assert f"{comprehensive_report_state.risk_score:.1f}/10" in md_text

        # 2. Milestones Table
        assert "Tiến Độ Các Cột Mốc Chiến Lược" in md_text
        assert "| **Trinh Sát Bề Mặt** |" in md_text
        assert "| **Rà Soát Lỗ Hổng** |" in md_text
        assert "| **Xác Minh Khai Thác** |" in md_text

        # 3. Severity Distribution
        assert "| Mức độ | CRITICAL | HIGH | MEDIUM | LOW | INFO | TỔNG CỘNG |" in md_text

        # 4. Attack Surface Tables
        assert "### 3.1. Các Cổng Dịch Vụ Mở & Dịch Vụ Lắng Nghe" in md_text
        assert "`8080/tcp`" in md_text
        assert "### 3.4. Tập Tin & Đường Dẫn Nhạy Cảm Lộ Lọt" in md_text
        assert "`/.env`" in md_text
        assert "### 3.5. Lỗ Hổng Cấu Hình CORS & Tiêu Đề Bảo Mật" in md_text
        assert "### 3.7. Chứng Chỉ SSL/TLS & Mã Hóa Giao Thức" in md_text
        assert "### 3.8. An Toàn Hạ Tầng DNS & Chống Giả Mạo Email" in md_text
        assert "### 3.10. An Toàn Cookie & Quản Lý Phiên" in md_text

        # 5. Compound Attack Scenarios
        assert "## 4. KỊCH BẢN KHAI THÁC PHỨC HỢP" in md_text

        # 6. Kill Chain Methodology Timeline
        assert "## 5. NHẬT KÝ THỰC THI KILL CHAIN" in md_text
        assert "`docker_scan_ports_fast`" in md_text
        assert "`docker_nuclei_scan`" in md_text

        # 7. Detailed Findings with code blocks
        assert "## 6. DANH MỤC PHÁT HIỆN & HỒ SƠ LỖ HỔNG CHI TIẾT" in md_text
        assert "CVE-2021-41773" in md_text
        assert "```text" in md_text
        assert "uid=0(root)" in md_text

        # 8. Prioritized Remediation Roadmap
        assert "## 7. LỘ TRÌNH KHẮC PHỤC PHÂN TẦNG THỜI GIAN" in md_text
        assert "Khẩn cấp (< 24 giờ)" in md_text
        assert "Ngắn hạn (< 7 ngày)" in md_text
        assert "Trung hạn (< 30 ngày)" in md_text

        # 9. Technical Appendix
        assert "## 8. PHỤ LỤC KỸ THUẬT THỰC THI" in md_text


class TestJsonReportCompleteness:
    """Verify JSON export contains 100% structured data."""

    def test_json_to_dict_completeness(self, comprehensive_report_state):
        data = comprehensive_report_state.to_dict()

        expected_keys = [
            "target", "scan_mode", "mission_objective", "start_time", "end_time",
            "duration", "executive_summary", "findings", "compound_threats",
            "methodology", "risk_score", "risk_label", "severity_counts",
            "recommendations", "prioritized_roadmap", "conclusion",
            "total_iterations", "red_teamer_final_answer", "attack_surface",
            "goal_progress", "milestones",
        ]

        for k in expected_keys:
            assert k in data, f"Key '{k}' missing from to_dict()"

        # Verify prioritized roadmap contents
        roadmap = data["prioritized_roadmap"]
        assert len(roadmap) >= 3
        phases = [r["phase"] for r in roadmap]
        assert any("Khẩn cấp" in p for p in phases)
        assert any("Ngắn hạn" in p for p in phases)
        assert any("Trung hạn" in p for p in phases)

        # Verify compound threats
        assert len(data["compound_threats"]) >= 1

        # Verify attack surface keys
        surface = data["attack_surface"]
        assert "open_ports" in surface
        assert "detected_technologies" in surface
        assert "parameterized_endpoints" in surface
        assert "exposed_sensitive_files" in surface
        assert "cors_issues" in surface
        assert "ssl_cert_info" in surface
        assert "dns_security_info" in surface
        assert "cookie_issues" in surface
